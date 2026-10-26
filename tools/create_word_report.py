"""Genera el informe editable de Word con evidencias aprobadas y capturas reales."""
from pathlib import Path
from xml.etree import ElementTree as ET
import json
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "output/pruebas"
TARGET = ROOT / "output/word/Pruebas_Entre_Paginas.docx"
REPO = "https://github.com/Anderson092205/Desafio3_DSM"
CAPTIONS = {
"00": ("Inicio de sesión", "Pantalla de acceso a la biblioteca mediante correo electrónico y contraseña."),
"01": ("Validación de acceso", "El correo incorrecto y una contraseña menor de seis caracteres muestran errores antes de enviar datos a Firebase."),
"02": ("Confirmación de contraseña", "El registro se bloquea si las contraseñas no coinciden. La contraseña y su confirmación se ocultan de forma predeterminada."),
"03": ("Registro completado", "La nueva cuenta accede a su inicio y puede definir su propia meta de lectura."),
"04": ("Meta inválida", "La meta de cero libros se rechaza. El mensaje indica que debe ser un número entero mayor que cero."),
"05": ("Meta anual guardada", "La meta de cinco libros queda guardada en Realtime Database para la cuenta autenticada."),
"06": ("Validación de libro", "El formulario identifica el título, el autor y el total de páginas inválidos."),
"07": ("Biblioteca y progreso", "El principito aparece como Leyendo. Sus 50 de 200 páginas equivalen al 25 por ciento de avance."),
"08": ("Lectura en el inicio", "La lectura en curso aparece junto a la meta anual y su porcentaje de avance."),
"09": ("Calificación obligatoria", "El formulario impide guardar un libro terminado sin una calificación de una a cinco estrellas."),
"10": ("Finalización de lectura", "El formulario permite asignar cinco estrellas al libro terminado."),
"11": ("Libro terminado", "La biblioteca muestra el estado Terminado, el 100 por ciento de avance, cinco estrellas y la fecha de finalización."),
"12": ("Estadísticas anuales", "El año actual contiene un libro terminado y 200 páginas. El conteo por género identifica Novela."),
"13": ("Libro mejor calificado", "La estadística identifica la lectura mejor calificada. En un empate se utiliza la fecha de finalización más temprana."),
"14": ("Avance de la meta", "El inicio refleja un libro terminado de una meta de cinco y cuatro libros pendientes para cumplirla."),
"15": ("Sesión cerrada", "Después de cerrar sesión se vuelve a la pantalla de acceso."),
"16": ("Biblioteca de otra cuenta", "La segunda cuenta tiene una biblioteca vacía. El servidor también impide que lea los datos de la primera."),
"17": ("Datos recuperados al ingresar", "Al ingresar otra vez con la primera cuenta, se recupera el libro terminado y su calificación."),
"18": ("Sesión después de cerrar el proceso", "Tras el cierre forzado del proceso y la reapertura, la cuenta continúa autenticada. No se reinició el teléfono."),
"19": ("Datos después de reabrir la aplicación", "La biblioteca conserva el libro terminado, el avance y la calificación. La meta también se recuperó al reabrir el proceso.")
}

def load(name):
    return json.loads((EVIDENCE / name).read_text(encoding="utf-8-sig"))

def unit_results():
    result = dict(tests=0, failures=0, errors=0, skipped=0)
    for path in (ROOT / "app/build/test-results/testDebugUnitTest").glob("TEST-*.xml"):
        suite = ET.parse(path).getroot()
        for key in result:
            result[key] += int(suite.attrib.get(key, 0))
    return result

def run_font(run, size=None, bold=None, color="000000"):
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    names = rpr.find(qn("w:rFonts"))
    if names is None:
        names = OxmlElement("w:rFonts"); rpr.append(names)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        names.set(qn("w:" + key), "Calibri")
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold

def text(doc, value, size=11, style=None):
    p = doc.add_paragraph(style=style)
    run_font(p.add_run(value), size=size)
    return p

def heading(doc, value, level=1):
    p = doc.add_heading(value, level)
    for run in p.runs: run_font(run, bold=True)
    p.paragraph_format.keep_with_next = True
    return p

def hyperlink(p, label, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "000000"); props.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
    run.append(props); value = OxmlElement("w:t"); value.text = label; run.append(value)
    link.append(run); p._p.append(link)

def table(doc, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(widths)); t.autofit = False
    for column, width in zip(t.columns, widths): column.width = Inches(width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            border.set(qn("w:" + key), value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for i, values in enumerate(rows):
        row = t.rows[0] if i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0: trpr.append(OxmlElement("w:tblHeader"))
        for j, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side, amount in (("top", 75), ("bottom", 75), ("left", 90), ("right", 90)):
                item = OxmlElement("w:" + side)
                item.set(qn("w:w"), str(amount)); item.set(qn("w:type"), "dxa"); mar.append(item)
            props.append(mar)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "404040" if i == 0 else ("F2F2F2" if i % 2 == 0 else "FFFFFF"))
            props.append(shade)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == len(widths) - 1 else WD_ALIGN_PARAGRAPH.LEFT
            run_font(p.add_run(str(value)), size=size, bold=i == 0, color="FFFFFF" if i == 0 else "000000")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def section_setup(s, columns=1):
    s.page_width = Inches(8.2677); s.page_height = Inches(11.6929)
    s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(.65)
    s.footer_distance = Inches(.28); s.header_distance = Inches(.25)
    cols = s._sectPr.find(qn("w:cols"))
    if cols is None: cols = OxmlElement("w:cols"); s._sectPr.append(cols)
    for child in list(cols): cols.remove(child)
    cols.set(qn("w:num"), str(columns)); cols.set(qn("w:space"), "360"); cols.set(qn("w:equalWidth"), "1")

def capture(doc, path, number, column_width):
    title, detail = CAPTIONS[path.stem.split("_", 1)[0]]
    p = text(doc, f"Figura {number} {title}", 12)
    p.runs[0].bold = True; p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(7)
    with Image.open(path) as image: width_px, height_px = image.size
    width = min(column_width - .06, 7.0 * width_px / height_px)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7); p.paragraph_format.keep_with_next = True
    p.paragraph_format.line_spacing = 1
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", title + ". " + detail); shape._inline.docPr.set("title", title)
    p = text(doc, detail, 10)
    p.paragraph_format.space_after = Pt(5); p.paragraph_format.line_spacing = 1.05

def build_report():
    backend, ui, process, device, build = [load(name) for name in (
        "firebase-results.json", "resultado_pruebas.json", "persistencia_proceso.json", "dispositivo.json", "compilacion.json")]
    units = unit_results()
    assert ui["resultado"] == process["resultado"] == build["compilation"] == "APROBADO"
    assert units["tests"] > 0 and units["failures"] == units["errors"] == 0
    assert all(item["passed"] is True for item in backend["results"])
    assert all(process[key] is True for key in ("uidRestaurado", "bibliotecaRecuperada", "metaRecuperada"))
    paths = [EVIDENCE / "capturas" / name for name in ui["capturas"] + process["capturas"]]
    assert len(paths) == 20 and all(path.is_file() for path in paths)
    doc = Document(); section_setup(doc.sections[0])
    doc.core_properties.title = "Pruebas de funcionamiento de Entre páginas"
    doc.core_properties.subject = "Desafío práctico 3 de Desarrollo de Software para Móviles"
    doc.core_properties.author = ""; doc.core_properties.last_modified_by = ""
    for name, size, bold in (("Normal", 11, False), ("Title", 26, False), ("Subtitle", 12, False),
                             ("Heading 1", 16, True), ("Heading 2", 12, True), ("Caption", 10, False)):
        style = doc.styles[name]
        style.font.name = "Calibri"; style.font.size = Pt(size); style.font.bold = bold; style.font.color.rgb = RGBColor(0,0,0)
        color = style.element.get_or_add_rPr().find(qn("w:color"))
        if color is not None:
            for key in ("themeColor", "themeShade", "themeTint"): color.attrib.pop(qn("w:" + key), None)
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            border = ppr.find(qn("w:pBdr"))
            if border is not None: ppr.remove(border)
        style.paragraph_format.space_after = Pt(7); style.paragraph_format.line_spacing = 1.08
    doc.styles["Title"].paragraph_format.space_after = Pt(15)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)
    footer = doc.sections[0].footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(footer.add_run("Página "), 9)
    fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r"); value = OxmlElement("w:t"); value.text = "1"; r.append(value); fld.append(r); footer._p.append(fld)

    text(doc, "Pruebas de funcionamiento de Entre páginas", 26, "Title")
    text(doc, "Desafío práctico 3 de Desarrollo de Software para Móviles\n24 de octubre de 2026")
    text(doc, "La aplicación permite llevar una biblioteca personal, registrar el progreso de lectura y comparar los libros terminados con una meta anual. Las pruebas confirmaron el registro y acceso a las cuentas, el guardado de libros y metas, las estadísticas y el aislamiento entre usuarios.")
    heading(doc, "Resumen de resultados")
    table(doc, [["Prueba", "Evidencia", "Resultado"],
        ["Lógica y validaciones", f"{units['tests']} pruebas JUnit", f"{units['tests']} aprobadas"],
        ["Firebase real", f"{len(backend['results'])} comprobaciones del servidor", "Todas aprobadas"],
        ["Interfaz Android", f"{len(ui['verificaciones'])} verificaciones y {len(ui['capturas'])} capturas", "Aprobado"],
        ["Cierre del proceso", "Sesión, biblioteca y meta recuperadas; 2 capturas", "Aprobado"],
        ["Compilación", f"APK generado; {build['lintErrors']} errores y {build['lintWarnings']} advertencias de lint", "Aprobado"]],
        [1.45,4.1,1.4177], 9.5)
    heading(doc, "Entorno de las pruebas", 2)
    text(doc, f"Teléfono físico {device['fabricante']} {device['modelo']} con Android {device['android']} y API {device['api']}, pantalla de {device['resolucion']} píxeles y conexión ADB. Se utilizó Firebase remoto real en el proyecto {backend['projectId']}.", 10)
    text(doc, f"Versiones conservadas del proyecto: Android Gradle Plugin {build['androidGradlePlugin']}, compile SDK {build['compileSdk']}, target SDK {build['targetSdk']} y min SDK {build['minSdk']}. Firebase BoM {build['firebaseBom']} es compatible con el Kotlin integrado en la versión original de Android Gradle Plugin.", 10)
    heading(doc, "Repositorio del proyecto", 2)
    hyperlink(doc.add_paragraph(), REPO, REPO)
    text(doc, "El enlace corresponde al repositorio configurado. Los cambios de esta entrega están en la carpeta local; este informe no afirma que ya se hayan publicado en GitHub.", 10)
    heading(doc, "Criterio de páginas anuales", 2)
    text(doc, "La estadística anual suma las páginas de los libros terminados en ese año. El avance parcial se muestra en la biblioteca. La aplicación no almacena un historial diario de páginas.", 10)

    doc.add_page_break()
    heading(doc, "Resultados de Firebase")
    text(doc, "Estas comprobaciones se ejecutaron contra Authentication y Realtime Database. Un rechazo cuenta como aprobado cuando coincide con la validación o restricción de acceso esperada.", 10)
    rows = [["Comprobación", "Resultado observado", "Estado"]]
    rows += [[item["name"],item["observed"],"Aprobado" if item["passed"] else "Falló"] for item in backend["results"]]
    table(doc, rows, [4.0,1.75,1.2177])
    text(doc, "Las cuentas temporales y los datos creados por estas pruebas del servidor fueron eliminados al finalizar, de acuerdo con el registro de limpieza.", 10)

    doc.add_page_break()
    heading(doc, "Recorrido de la aplicación")
    text(doc, "Las pruebas de interfaz en el teléfono físico verificaron las siguientes operaciones. Las capturas de las páginas siguientes documentan distintas ejecuciones aprobadas de las pruebas, incluida la comprobación de cierre y reapertura del proceso.", 10)
    for i, check in enumerate(ui["verificaciones"], 1):
        p = text(doc, f"{i}. {check}", 10); p.paragraph_format.space_after = Pt(5)
    heading(doc, "Pruebas unitarias", 2)
    text(doc, f"JUnit registró {units['tests']} pruebas, {units['failures']} fallos, {units['errors']} errores y {units['skipped']} omitidas. Los casos cubren validaciones, porcentajes, límites numéricos, selección del año, zonas horarias y empates de estadísticas.", 10)
    heading(doc, "Persistencia después de cerrar el proceso", 2)
    text(doc, "Se recreó la actividad durante el recorrido. En una comprobación adicional se cerró forzadamente el proceso mediante ADB y se volvió a abrir la aplicación. Se verificaron la misma cuenta autenticada, la biblioteca con su calificación y la meta anual. Ambas comprobaciones fueron aprobadas. No se reinició el teléfono.", 10)
    heading(doc, "Archivos de evidencia", 2)
    text(doc, "Los resultados se conservan en output/pruebas y los resultados unitarios en app/build/test-results/testDebugUnitTest. El informe no incluye tokens, contraseñas ni identificadores internos de las cuentas en su texto.", 10)
    for start in range(0,len(paths),2):
        section = doc.add_section(WD_SECTION_START.NEW_PAGE); section_setup(section,2)
        page_width = (section.page_width-section.left_margin-section.right_margin)/914400
        column_width = (page_width-.25)/2
        capture(doc,paths[start],start+1,column_width)
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
        p.add_run().add_break(WD_BREAK.COLUMN)
        capture(doc,paths[start+1],start+2,column_width)
    TARGET.parent.mkdir(parents=True,exist_ok=True); doc.save(TARGET)
    check = Document(TARGET)
    assert len(check.inline_shapes) == 20 and len(check.tables) == 2 and len(check.sections) == 11
    assert check.paragraphs[0].style.name == "Title"
    print(json.dumps({"docx":str(TARGET),"images":20,"editableTables":2,"sections":11,"plannedPages":13,
        "unitTests":units["tests"],"firebaseChecks":len(backend["results"]),"uiChecks":len(ui["verificaciones"]),
        "renderStatus":"Pendiente de renderizado nativo y revisión visual por el agente principal"}, ensure_ascii=False,indent=2))

if __name__ == "__main__":
    build_report()
